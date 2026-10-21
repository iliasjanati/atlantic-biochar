"""Télécharge le jeu officiel « Production végétale 2010-2022 » de data.gov.ma.

Le fichier source contient des données de production/superficie agricole,
pas des tonnages directement mesurés de résidus. Il est conservé tel quel
dans data/raw afin de garder une trace de la source et de sa licence.

Usage depuis la racine du projet :
    python scripts/collecte_open_data_maroc.py
    python scripts/collecte_open_data_maroc.py --output-dir data/raw
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


API_URL = "https://data.gov.ma/data/api/3/action/package_search"
DATASET_TITLE = "Production végétale 2010-2022"
DATASET_PAGE = "https://data.gov.ma/data/dataset/api_dataset/production-vegetale-2010-2022"
USER_AGENT = "AtlanticBiocharDataCollector/1.0 (open-data research)"


def fetch_json(url: str) -> dict:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=45) as response:
        return json.load(response)


def find_dataset() -> dict:
    query = urllib.parse.urlencode({"q": 'title:"Production végétale 2010-2022"', "rows": 20})
    response = fetch_json(f"{API_URL}?{query}")
    if not response.get("success"):
        raise RuntimeError("L’API du portail a répondu sans succès.")

    results = response.get("result", {}).get("results", [])
    exact = [item for item in results if item.get("title", "").strip().casefold() == DATASET_TITLE.casefold()]
    if not exact:
        raise RuntimeError(
            f"Jeu de données introuvable dans l’API. Consultez {DATASET_PAGE} et réessayez plus tard."
        )
    return exact[0]


def choose_xlsx(dataset: dict) -> dict:
    resources = dataset.get("resources", [])
    candidates = [
        resource
        for resource in resources
        if "xlsx" in str(resource.get("format", "")).casefold()
        or str(resource.get("url", "")).split("?", 1)[0].casefold().endswith(".xlsx")
    ]
    if not candidates:
        raise RuntimeError("Aucune ressource Excel (XLSX) n’est publiée pour ce jeu de données.")
    return candidates[0]


def safe_filename(name: str) -> str:
    name = re.sub(r"[^\w.-]+", "_", name, flags=re.UNICODE).strip("_.")
    return name or "production_vegetale_2010_2022.xlsx"


def download(url: str, destination: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=90) as response, destination.open("wb") as output:
        while True:
            chunk = response.read(1024 * 1024)
            if not chunk:
                break
            output.write(chunk)


def normalize_xlsx(source: Path, destination: Path) -> int:
    """Met production et superficie sur une ligne par campagne et produit."""
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip().casefold() for value in next(rows)]
    required = {"occurrence", "filière", "produit", "indicateur", "valeur"}
    if not required.issubset(headers):
        raise RuntimeError(f"Colonnes inattendues dans le fichier Excel : {headers}")
    indexes = {name: headers.index(name) for name in required}
    values: dict[tuple[str, str, str], dict[str, float | None]] = {}

    for row in rows:
        if not row or all(value is None for value in row):
            continue
        campaign = str(row[indexes["occurrence"]] or "").strip()
        sector = str(row[indexes["filière"]] or "").strip()
        product = str(row[indexes["produit"]] or "").strip()
        indicator = str(row[indexes["indicateur"]] or "").strip().casefold()
        raw_value = row[indexes["valeur"]]
        try:
            numeric_value = float(raw_value) if raw_value is not None else None
        except (TypeError, ValueError):
            numeric_value = None
        key = (campaign, sector, product)
        item = values.setdefault(key, {"production_t": None, "superficie_ha": None})
        if "production" in indicator and ("tonne" in indicator or "-t" in indicator):
            item["production_t"] = numeric_value
        elif "superficie" in indicator:
            item["superficie_ha"] = numeric_value

    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.DictWriter(
            output,
            fieldnames=["campagne", "filiere", "produit", "production_t", "superficie_ha", "rendement_t_ha"],
            delimiter=";",
        )
        writer.writeheader()
        for (campaign, sector, product), item in sorted(values.items()):
            production = item["production_t"]
            area = item["superficie_ha"]
            writer.writerow(
                {
                    "campagne": campaign,
                    "filiere": sector,
                    "produit": product,
                    "production_t": production if production is not None else "",
                    "superficie_ha": area if area is not None else "",
                    "rendement_t_ha": round(production / area, 6) if production is not None and area else "",
                }
            )
    workbook.close()
    return len(values)


def main() -> int:
    parser = argparse.ArgumentParser(description="Télécharge les données agricoles officielles de data.gov.ma.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("data/raw"),
        help="Dossier de destination, relatif à l’emplacement depuis lequel le script est lancé (défaut : data/raw).",
    )
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()

    try:
        dataset = find_dataset()
        resource = choose_xlsx(dataset)
        output_dir.mkdir(parents=True, exist_ok=True)
        original_name = resource.get("name") or resource.get("url", "").rsplit("/", 1)[-1]
        filename = safe_filename(original_name)
        if not filename.casefold().endswith(".xlsx"):
            filename += ".xlsx"
        data_path = output_dir / filename

        print(f"Téléchargement : {dataset.get('title', DATASET_TITLE)}")
        download(resource["url"], data_path)
        normalized_path = output_dir / "production_vegetale_maroc_normalisee.csv"
        normalized_rows = normalize_xlsx(data_path, normalized_path)

        metadata = {
            "titre": dataset.get("title", DATASET_TITLE),
            "producteur": dataset.get("organization", {}).get("title"),
            "licence": dataset.get("license_title"),
            "page_source": DATASET_PAGE,
            "api_source": API_URL,
            "ressource_nom": resource.get("name"),
            "ressource_url": resource.get("url"),
            "fichier_local": data_path.name,
            "fichier_normalise": normalized_path.name,
            "nombre_lignes_normalisees": normalized_rows,
            "date_collecte_utc": datetime.now(timezone.utc).isoformat(),
            "limite": "Production et superficies agricoles; les tonnages de résidus doivent être estimés séparément.",
        }
        metadata_path = data_path.with_suffix(".metadata.json")
        metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")

        print(f"Fichier enregistré : {data_path}")
        print(f"Tableau nettoyé enregistré : {normalized_path} ({normalized_rows} lignes)")
        print(f"Métadonnées enregistrées : {metadata_path}")
        print("Attention : ce fichier ne fournit pas directement les tonnages de résidus.")
        return 0
    except (urllib.error.URLError, TimeoutError, json.JSONDecodeError, KeyError, RuntimeError) as error:
        print(f"Échec de la collecte : {error}", file=sys.stderr)
        print("Vérifiez la connexion et la disponibilité du portail data.gov.ma.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
